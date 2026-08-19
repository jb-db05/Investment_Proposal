#!/usr/bin/env python3
"""
Rebuild the "proposed portfolio" line-items pages as real PowerPoint tables.

The reference deck (work16.pptx) built these two slides out of tall
textboxes with one paragraph per row, manually positioned to line up like a
table. It breaks the moment a category has more lines than fit the fixed
paragraph count it was built for, and it can't repaginate. This script
replaces that with actual `a:tbl` tables (python-pptx `add_table`), grouped
by category, colored exactly like the reference deck's own category bars,
with a slide count that adapts to however many holdings `parsed.json` has.

Layout rule (matches the reference deck exactly): each top-level asset
class gets ONE colored header row carrying both its name (left) and its
total weight (right) on the same line — never a separate subtotal row.
Finer sub-groups (a fixed-income bucket, an equity region) show their own
label in the merged first column but get no subtotal of their own; that
detail lives on the sleeve-breakdown slides instead.

Hard rule (matches the task spec): a sub-group ("Govies 1-10 (local)", "US"
equities, "Cash") never splits across two slides. If it doesn't fit on the
current slide, the whole group moves to the next one.

Usage (also called from build_proposal.py):
    python build_line_items_tables.py deck.pptx parsed.json -o deck_out.pptx
"""
from __future__ import annotations

import argparse
import json
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# --- visual constants, matched to the template's own title-block geometry ---
TABLE_LEFT = Emu(665163)
TABLE_WIDTH = Emu(9398001)
TABLE_TOP = Emu(2140000)
TABLE_BOTTOM_MARGIN = Emu(520000)  # keep clear of the footer/page-number

COL_WIDTHS = [Emu(1900000), Emu(1500000), Emu(4573001), Emu(1425000)]  # category, ISIN, instrument, weight
ROW_HEIGHT = Emu(205000)  # tightened from 228600 to fit more rows per slide

# Category bar colors. Cash/Fixed Income/Private Assets/Commodities/Total
# are extracted directly from the reference deck's own Rectangle shapes
# (slides 9-10) via their theme scheme references: Cash=bg2->lt2, Fixed
# Income=accent6, Alternatives=accent5, Commodities=accent2, Total=tx2->
# dk2. Equities and Structured Products don't exist as their own colored
# categories in the reference deck's own example (Equities originally
# shared Fixed Income's accent6) -- per explicit design direction they use
# Sky Blue and Peach Pink instead, both distinct from every other bar.
CATEGORY_COLORS = {
    "Cash": RGBColor(0xFF, 0xC5, 0x45),
    "Fixed Income": RGBColor(0x4B, 0x5F, 0x80),
    "Equities": RGBColor(0x79, 0xD6, 0xFF),          # Sky Blue
    "Private Assets": RGBColor(0xAC, 0x5D, 0x85),
    "Commodities": RGBColor(0x3B, 0xAF, 0x90),
    "Structured Products": RGBColor(0xE6, 0xA4, 0xAD),  # Peach Pink
    "Other": RGBColor(0xFF, 0xA4, 0x00),
}
TOTAL_COLOR = RGBColor(0x20, 0x29, 0x45)
NAVY = RGBColor(0x20, 0x29, 0x45)
HEADER_TEXT_COLOR = RGBColor(0xFF, 0xFF, 0xFF)
FONT_NAME = "Nunito Sans ExtraBold"
BODY_FONT_NAME = "Calibri"


def category_color(top_category: str) -> RGBColor:
    return CATEGORY_COLORS.get(top_category, RGBColor(0x4B, 0x5F, 0x80))


@dataclass
class Group:
    top_category: str
    group: str
    rows: list[dict]


def group_line_items(line_items: list[dict]) -> list[Group]:
    groups: list[Group] = []
    for li in line_items:
        if groups and groups[-1].top_category == li["top_category"] and groups[-1].group == li["group"]:
            groups[-1].rows.append(li)
        else:
            groups.append(Group(li["top_category"], li["group"], [li]))
    return groups


def paginate(groups: list[Group], rows_per_slide: int) -> list[list[Group]]:
    """Greedy bin-packing where a Group (fine sub-bucket) is the atomic,
    never-split unit. A category-header row is only counted once per
    top_category per page (repeated on a continuation page if that
    category's groups spill over)."""
    pages: list[list[Group]] = []
    current: list[Group] = []
    current_rows = 0
    last_top_category = None

    def group_cost(g: Group, top_cat_changed: bool) -> int:
        return (1 if top_cat_changed else 0) + len(g.rows)  # header? + data rows only, no subtotal row

    for g in groups:
        top_cat_changed = g.top_category != last_top_category
        cost = group_cost(g, top_cat_changed)
        if current and current_rows + cost > rows_per_slide:
            pages.append(current)
            current = []
            current_rows = 0
            last_top_category = None
            top_cat_changed = True
            cost = group_cost(g, top_cat_changed)
        current.append(g)
        current_rows += cost
        last_top_category = g.top_category
    if current:
        pages.append(current)
    return pages


def _set_cell(cell, text, *, bold=False, fill=None, font_color=NAVY, size=10, align=PP_ALIGN.LEFT, font_name=BODY_FONT_NAME):
    cell.margin_left = Emu(45720)
    cell.margin_right = Emu(45720)
    cell.margin_top = Emu(9144)
    cell.margin_bottom = Emu(9144)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run() if not p.runs else p.runs[0]
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font_name
    run.font.color.rgb = font_color


def _merge_col1(table, r0, r1):
    if r1 > r0:
        table.cell(r0, 0).merge(table.cell(r1, 0))


def render_page(slide, page_groups: list[Group], n_rows: int, is_last_page: bool,
                 grand_total_pct: float | None, category_totals: dict[str, float]):
    n_cols = 4
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, TABLE_LEFT, TABLE_TOP,
                                            TABLE_WIDTH, ROW_HEIGHT * n_rows)
    table = graphic_frame.table
    table.first_row = False
    table.horz_banding = False
    for i, w in enumerate(COL_WIDTHS):
        table.columns[i].width = w
    for r in range(n_rows):
        table.rows[r].height = ROW_HEIGHT

    # column-label row: navy text on white, not a colored bar — the colored
    # bars are reserved for the asset-class header rows below
    headers = ["Asset Allocation", "ISIN", "Instrument", "Weight"]
    for c, h in enumerate(headers):
        _set_cell(table.cell(0, c), h, bold=True, fill=RGBColor(0xFF, 0xFF, 0xFF), font_color=NAVY,
                  size=10, font_name=FONT_NAME)
    r = 1
    last_top_category = None
    for g in page_groups:
        if g.top_category != last_top_category:
            color = category_color(g.top_category)
            total_pct = category_totals.get(g.top_category, 0.0)
            for c in range(n_cols):
                if c == 0:
                    text = g.top_category
                elif c == 3:
                    text = f"{total_pct:.1f}%"
                else:
                    text = ""
                _set_cell(table.cell(r, c), text, bold=True, fill=color, font_color=HEADER_TEXT_COLOR,
                          size=10.5, font_name=FONT_NAME, align=(PP_ALIGN.RIGHT if c == 3 else PP_ALIGN.LEFT))
            r += 1
            last_top_category = g.top_category
        group_first_row = r
        for li in g.rows:
            _set_cell(table.cell(r, 1), li["isin"] or "", size=9.5)
            _set_cell(table.cell(r, 2), li["instrument"] or "", size=9.5)
            _set_cell(table.cell(r, 3), f"{li['weight_pct']:.1f}%", size=9.5, align=PP_ALIGN.RIGHT)
            r += 1
        _set_cell(table.cell(group_first_row, 0), g.group, bold=True, size=9.5)
        _merge_col1(table, group_first_row, r - 1)

    if is_last_page and grand_total_pct is not None:
        for c in range(n_cols):
            text = "Total" if c == 0 else (f"{grand_total_pct:.1f}%" if c == 3 else "")
            _set_cell(table.cell(r, c), text, bold=True, fill=TOTAL_COLOR, font_color=HEADER_TEXT_COLOR, size=10,
                      font_name=FONT_NAME, align=(PP_ALIGN.RIGHT if c == 3 else PP_ALIGN.LEFT))
        r += 1
    assert r == n_rows, f"row count mismatch: built {r}, allocated {n_rows}"


def find_slide(prs, predicate):
    for i, slide in enumerate(prs.slides):
        if predicate(slide):
            return i, slide
    return None, None


def is_line_items_slide(slide) -> bool:
    names = {sh.name for sh in slide.shapes}
    return names == {"Title 1", "Text Placeholder 2"}


def duplicate_slide_after(prs, index):
    """Insert a new slide using the same layout as prs.slides[index],
    positioned immediately after it. Returns the new slide."""
    src = prs.slides[index]
    new_slide = prs.slides.add_slide(src.slide_layout)  # appended at the end
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    new_el = slides[-1]
    xml_slides.remove(new_el)
    xml_slides.insert(index + 1, new_el)
    return new_slide


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get(qn("r:id"))
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


# Same fix as build_proposal.py's set_subtitle(): the subtitle placeholder
# uses <a:spAutoFit/> with anchor="b" at the layout level, so a subtitle
# that wraps to a second line visibly shifts down instead of growing
# upward. Recreated as a plain, fixed-height, top-anchored textbox at the
# template's own position (H=1.85cm, V=4.8cm) instead of edited in place.
SUBTITLE_LEFT = Emu(665163)
SUBTITLE_TOP = Emu(1728947)
SUBTITLE_WIDTH = Emu(9398001)
SUBTITLE_HEIGHT = Emu(560000)


def set_subtitle(slide, text):
    old = None
    for sh in slide.shapes:
        if sh.name == "Text Placeholder 2":
            old = sh
            break
    if old is not None:
        old._element.getparent().remove(old._element)
    tb = slide.shapes.add_textbox(SUBTITLE_LEFT, SUBTITLE_TOP, SUBTITLE_WIDTH, SUBTITLE_HEIGHT)
    tb.name = "Text Placeholder 2"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    return tb


def set_placeholder_text(slide, name, text):
    for sh in slide.shapes:
        if sh.name == name and sh.has_text_frame:
            tf = sh.text_frame
            p = tf.paragraphs[0]
            for run in list(p.runs)[1:]:
                run._r.getparent().remove(run._r)
            if p.runs:
                p.runs[0].text = text
            else:
                p.add_run().text = text
            for extra_p in list(tf.paragraphs[1:]):
                extra_p._p.getparent().remove(extra_p._p)
            return


def build(prs: Presentation, parsed: dict) -> Presentation:
    slide_h = prs.slide_height
    usable = slide_h - TABLE_TOP - TABLE_BOTTOM_MARGIN
    rows_per_slide = max(8, int(usable // ROW_HEIGHT))

    line_items = parsed["line_items"]
    groups = group_line_items(line_items)
    pages = paginate(groups, rows_per_slide)
    n_pages = len(pages)
    grand_total = sum(li["weight_pct"] for li in line_items)
    # Category totals come straight from parsed.json's own asset-allocation
    # breakdown (already computed once, correctly) rather than resumming
    # line_items per page, so a category split across pages still shows its
    # one true total everywhere its header appears.
    category_totals = dict(parsed["asset_allocation_pct"])

    anchor_idx, anchor_slide = find_slide(prs, is_line_items_slide)
    if anchor_idx is None:
        raise RuntimeError("No blank line-items slide found in the template "
                            "(expected a slide with only 'Title 1' and 'Text Placeholder 2' shapes)")
    # collect every consecutive blank line-items slide starting at anchor (the
    # template ships two: the old 'cash and fixed income' / 'equity...' slots)
    slot_indices = [anchor_idx]
    probe = anchor_idx + 1
    while probe < len(prs.slides) and is_line_items_slide(prs.slides[probe]):
        slot_indices.append(probe)
        probe += 1

    while len(slot_indices) < n_pages:
        new_slide = duplicate_slide_after(prs, slot_indices[-1])
        new_idx = slot_indices[-1] + 1
        slot_indices.append(new_idx)
    while len(slot_indices) > n_pages:
        delete_slide(prs, slot_indices.pop())

    ccy = parsed["base_currency"]
    total_m = parsed["total_value_eur"] / 1_000_000
    date = parsed["valuation_date"]

    for page_num, (idx, page_groups) in enumerate(zip(slot_indices, pages), start=1):
        slide = prs.slides[idx]
        title = "Proposed portfolio: full holdings list" if n_pages > 1 else "Proposed portfolio"
        subtitle = (f"All positions of the proposed {ccy} {total_m:.1f}m allocation, "
                    f"part {page_num} of {n_pages}. Weights as at {date}.")
        set_placeholder_text(slide, "Title 1", title)
        set_subtitle(slide, subtitle)

        n_rows = 1  # column header
        last_tc = None
        for g in page_groups:
            if g.top_category != last_tc:
                n_rows += 1
                last_tc = g.top_category
            n_rows += len(g.rows)
        is_last = page_num == n_pages
        if is_last:
            n_rows += 1
        render_page(slide, page_groups, n_rows, is_last, grand_total if is_last else None, category_totals)

    return prs


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx_path")
    ap.add_argument("parsed_json")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()

    prs = Presentation(args.pptx_path)
    parsed = json.loads(open(args.parsed_json).read())
    build(prs, parsed)
    prs.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
