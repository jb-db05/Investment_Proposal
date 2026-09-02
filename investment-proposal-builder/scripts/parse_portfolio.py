#!/usr/bin/env python3
"""
Parse a client custodian Excel export into parsed.json — the single data
file every other build script (update_chart.py, build_line_items_tables.py,
build_proposal.py) reads from. Never edit parsed.json by hand; re-run this
script if the source Excel changes.

Every rule this script applies is documented in ../references/assumptions.md
by section number; the code comments below reference those sections instead
of re-explaining the rule.

Usage:
    python parse_portfolio.py Portfolio.xlsx --profile Balanced -o parsed.json
    python parse_portfolio.py Portfolio.xlsx --profile Balanced \
        --bucket-overrides overrides.csv --market-cap-overrides mcap.csv \
        -o parsed.json
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "references"))
from geography_map import region_of, equity_region_of  # noqa: E402

PORTFOLIO_SECTIONS = {
    "Cash",
    "Short-term instruments",
    "Bonds",
    "Equities",
    "Structured products",
    "Private Assets",
    "Gold and other commodities",
    "Other investments",
}

# Maps a raw Portfolio-sheet section header to the asset-class bucket used
# throughout parsed.json and the slide 9/10 line-items table (assumptions.md §4-5).
ASSET_CLASS_MAP = {
    "Cash": "Cash",
    "Short-term instruments": "Fixed Income",
    "Bonds": "Fixed Income",
    "Equities": "Equities",
    "Structured products": "Structured Products",
    "Private Assets": "Private Assets",
    "Gold and other commodities": "Commodities",
    "Other investments": "Other",
}

# ISO 4217 codes for the four precious metals. A "current account" denominated
# in one of these is a metal position held in account form, not a cash balance:
# it carries the metal's price risk, not a currency's (assumptions.md §1).
PRECIOUS_METAL_CURRENCIES = {"XAU", "XAG", "XPT", "XPD"}

GROWTH_CLASSES = {"Equities", "Alternatives", "Commodities", "Structured Products", "Private Assets"}
DEFENSIVE_CLASSES = {"Cash", "Fixed Income"}

RISK_PROFILES = ["Fixed Income", "Conservative", "Moderate", "Balanced", "Growth", "Equity"]

IG_RATINGS = {"AAA", "AA+", "AA", "AA-"}
CORP_IG_RATINGS = {"A+", "A", "A-", "BBB+", "BBB", "BBB-"}

FUND_KEYWORDS = ("FUND", " FD ", "SICAV", "ETF", "ETC")
STRUCTURED_KEYWORDS = ("CERTIFICATE", " CERT", "AMC", "TRACKER", "DCI-", "RAF ", "CPN ")


def is_precious_metal_account(description, currency) -> bool:
    """A cash-section balance denominated in gold, silver, platinum or
    palladium (assumptions.md §1). The currency code is what decides — the
    description is only checked to keep the rule to account balances, so a
    metal-denominated security would still be classified by its own section."""
    return (str(currency or "").strip().upper() in PRECIOUS_METAL_CURRENCIES
            and "ACCOUNT" in str(description or "").upper())


def norm_header(h) -> str:
    return " ".join(str(h or "").split()).strip().lower()


def find_header_row(ws) -> tuple[int | None, dict[str, int]]:
    """The custodian export's own header row: the first row with >=5 filled
    cells, one of which names an ISIN column. Returns its 1-based row index
    and a {normalized header name: column index} map, or (None, {}) if this
    sheet has no such row (i.e. it isn't a holdings export)."""
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        non_empty = [v for v in vals if v not in (None, "")]
        if len(non_empty) >= 5 and any(isinstance(v, str) and "isin" in v.lower() for v in vals):
            return r, {norm_header(v): c for c, v in enumerate(vals, 1) if v not in (None, "")}
    return None, {}


def pick_portfolio_sheet(wb):
    """The holdings sheet. The custodian names it 'Portfolio', but the same
    export re-saved out of Excel or forwarded by a client routinely arrives
    with a default sheet name ('Sheet1', 'Feuil1', an account number) — that
    is a naming accident, not a different file format, so fall back to the
    only sheet that carries a holdings header row rather than making the user
    rename their own file. The 'Fixed Income' proposal tab is excluded: it has
    its own parser (parse_proposed_bonds) and its own meaning."""
    if "Portfolio" in wb.sheetnames:
        return wb["Portfolio"]
    candidates = [ws for ws in wb.worksheets
                  if ws.title != "Fixed Income" and find_header_row(ws)[0] is not None]
    if not candidates:
        sys.exit("Error: no sheet in this file has a holdings header row (a row "
                 "with an 'ISIN code' column) — expected a 'Portfolio' sheet")
    if len(candidates) > 1:
        sys.exit("Error: more than one sheet looks like a holdings export "
                 f"({', '.join(ws.title for ws in candidates)}) — name the "
                 "right one 'Portfolio' so there is no ambiguity")
    print(f"Note: no 'Portfolio' sheet — using '{candidates[0].title}', the only "
          "sheet with a holdings header row.")
    return candidates[0]


def load_sheet_rows(ws) -> tuple[dict[str, int], list[dict]]:
    """Find the header row (see find_header_row) then yield every row below it
    as a dict keyed by normalized header name, tagging each with its section
    from column A."""
    header_row_idx, headers = find_header_row(ws)
    if header_row_idx is None:
        raise ValueError(f"Could not find a header row with an ISIN column in sheet '{ws.title}'")

    rows = []
    section = None
    for r in range(header_row_idx + 1, ws.max_row + 1):
        col_a = ws.cell(r, 1).value
        if isinstance(col_a, str) and col_a.strip() in PORTFOLIO_SECTIONS:
            section = col_a.strip()
            continue
        row = {h: ws.cell(r, c).value for h, c in headers.items()}
        if all(v in (None, "") for v in row.values()):
            continue
        row["_section"] = section
        rows.append(row)
    return headers, rows


def load_bucket_overrides(path: str | None) -> dict[str, str]:
    if not path:
        return {}
    out = {}
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            out[row["isin"].strip()] = row["bucket"].strip()
    return out


def load_market_cap_overrides(path: str | None) -> dict[str, str]:
    if not path:
        return {}
    out = {}
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            out[row["isin"].strip()] = row["market_cap"].strip()
    return out


def load_sector_overrides(path: str | None) -> dict[str, str]:
    if not path:
        return {}
    out = {}
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            out[row["isin"].strip()] = row["sector"].strip()
    return out


def classify_fi_bucket(row: dict, overrides: dict[str, str]) -> str:
    """assumptions.md §4. Note: the Portfolio sheet carries no Sector column
    (only the separate 'Fixed Income' proposal tab does), so this rule uses
    Composite Rating + Duration + Geographical breakdown only."""
    isin = str(row.get("isin code") or "").strip()
    if isin in overrides:
        return overrides[isin]

    geography = row.get("geographical breakdown")
    if region_of(geography) == "MEA EM":
        return "EM Debt"

    rating = str(row.get("composite rating") or "").strip()
    duration = row.get("duration")

    if rating in IG_RATINGS:
        try:
            d = float(duration)
        except (TypeError, ValueError):
            d = 0.0
        return "Govies 10+ (local)" if d > 10 else "Govies 1-10 (local)"
    if rating in CORP_IG_RATINGS:
        return "Corporate IG (local)"
    return "High Yield (local or global hdg)"


def classify_vehicle(description: str, rating, coupon, currency=None, asset_class=None) -> str:
    """assumptions.md §6"""
    d = (description or "").upper()
    if any(k in d for k in FUND_KEYWORDS):
        return "Fund"
    if any(k in d for k in STRUCTURED_KEYWORDS):
        return "Structured / AMC"
    # The blank-rating/blank-coupon signature below is a signature of
    # *securities* — collective vehicles that carry neither. Three kinds of
    # holding carry neither for an entirely different reason and are not
    # funds: a bank account, a metal account (the metal itself, held in
    # account form), and physical metal in the commodities sleeve (bars and
    # coins — GOLD KG, OR KRUGERRAND). Funds and ETCs in that sleeve are
    # already caught by the keyword rule above, so what is left there is
    # physical: a direct line.
    if asset_class in ("Cash", "Commodities") or is_precious_metal_account(description, currency):
        return "Direct line"
    if (rating in (None, "", "Not Rated")) and coupon in (None, ""):
        return "Fund"
    return "Direct line"


def vehicle_of(row: dict) -> str:
    """classify_vehicle() for a parsed row (which already knows its asset class)."""
    return classify_vehicle(row.get("description"), row.get("composite rating"),
                            row.get("coupon (%)"), row.get("currency"), row.get("_asset_class"))


ILLIQUID_BUCKET = "Illiquid (lock-up)"


def classify_liquidity(row: dict, asset_class: str, vehicle: str) -> str:
    """assumptions.md §10. Five buckets, by how a holding is actually
    realised — not by its asset class and not by its vehicle alone."""
    if asset_class == "Private Assets" or "COMMIT" in str(row.get("description") or "").upper():
        return ILLIQUID_BUCKET
    # Account balances settle like money, whether the account is denominated
    # in a currency or in a metal (§1a) — a metal account is the currency form
    # of the metal, not a bar in a vault.
    if asset_class == "Cash" or is_precious_metal_account(row.get("description"), row.get("currency")):
        return "Cash & metal accounts"
    # What is left holding metal directly in the commodities sleeve is
    # physical: sellable at a screen price, but neither a listed security nor
    # a fund with a daily NAV.
    if asset_class == "Commodities" and vehicle == "Direct line":
        return "Physical assets"
    if vehicle == "Fund":
        return "Daily-liquid fund"
    return "Listed"


def describe_bond_subsleeve(row: dict) -> str:
    """Cosmetic label only, assumptions.md §11."""
    desc = str(row.get("description") or "")
    sector = str(row.get("sector") or "").strip()
    du = desc.upper()
    keyword_map = [
        ("AUTO", "Corporate bond (auto)"),
        ("OIL", "Corporate bond (oil)"),
        ("PORT", "Corporate bond (port)"),
        ("AIRCRAFT", "Corporate bond (aircraft)"),
        ("AIR ", "Corporate bond (aircraft)"),
        ("COMMOD", "Corporate bond (commodity)"),
        ("GAMING", "Corporate bond (gaming)"),
        ("CASINO", "Corporate bond (gaming)"),
        ("TELE", "Corporate bond (telecom)"),
        ("REIT", "Corporate bond (REIT)"),
        ("BANK", "Bank bond"),
        ("INSUR", "Insurance-linked security"),
        ("ILS ", "Insurance-linked security"),
    ]
    for kw, label in keyword_map:
        if kw in du:
            return label
    is_fund = any(k in du for k in FUND_KEYWORDS)
    if is_fund:
        return "High yield bond fund"
    if "SUPRANATIONAL" in sector.upper() or "DEVELOPMENT" in du:
        return "Supranational bond"
    if any(k in du for k in STRUCTURED_KEYWORDS):
        return "Structured note / certificate"
    if sector.lower() in ("government", ""):
        return "Quasi-sovereign bond"
    return "Corporate bond"


def maturity_bucket(maturity_val, as_of=None) -> str:
    """Bucket a maturity date into the 0-3 / 3-5 / 5-7 / 7-10 / 10+ years bands
    used on the 'Proposed bond selection' slide (slide_recipe.md)."""
    import datetime as _dt

    as_of = as_of or _dt.date.today()
    if isinstance(maturity_val, _dt.datetime):
        maturity = maturity_val.date()
    elif isinstance(maturity_val, _dt.date):
        maturity = maturity_val
    elif isinstance(maturity_val, str) and maturity_val.strip():
        maturity = None
        for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
            try:
                maturity = _dt.datetime.strptime(maturity_val.strip(), fmt).date()
                break
            except ValueError:
                continue
        if maturity is None:
            return "Unknown"
    else:
        return "Unknown"
    years = (maturity - as_of).days / 365.25
    if years <= 3:
        return "0 - 3 years"
    if years <= 5:
        return "3 - 5 years"
    if years <= 7:
        return "5 - 7 years"
    if years <= 10:
        return "7 - 10 years"
    return "10+ years"


def parse_proposed_bonds(wb) -> dict:
    """The 'Fixed Income' tab is a curated bond *proposal*, not part of the
    client's current holdings (assumptions.md §1) — every row equally
    weighted, matching the reference deck's 'N issues, equally weighted at
    X each' framing on the Proposed Bond Selection slide."""
    ws = wb["Fixed Income"]
    headers, rows_dict = {}, []
    header_row_idx = None
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(isinstance(v, str) and "isin" in v.lower() for v in vals):
            header_row_idx = r
            for c, v in enumerate(vals, 1):
                if v not in (None, ""):
                    headers[norm_header(v)] = c
            break
    if header_row_idx is None:
        return None
    for r in range(header_row_idx + 1, ws.max_row + 1):
        row = {h: ws.cell(r, c).value for h, c in headers.items()}
        if all(v in (None, "") for v in row.values()):
            continue
        rows_dict.append(row)

    n = len(rows_dict)
    if n == 0:
        return None
    equal_w = 100.0 / n
    amount_each = rows_dict[0].get("amount to invest")

    def bucket_sum(key_fn):
        d = defaultdict(float)
        for row in rows_dict:
            d[key_fn(row) or "Unknown"] += equal_w
        return {k: round(v, 4) for k, v in sorted(d.items(), key=lambda kv: -kv[1])}

    return {
        "num_issues": n,
        "amount_each": amount_each,
        "currency": rows_dict[0].get("currency"),
        "total_amount": amount_each * n if isinstance(amount_each, (int, float)) else None,
        "by_geography_pct": bucket_sum(lambda row: row.get("country")),
        "by_currency_pct": bucket_sum(lambda row: row.get("currency")),
        "by_maturity_pct": bucket_sum(lambda row: maturity_bucket(row.get("maturity"))),
        "by_rating_pct": bucket_sum(lambda row: row.get("rating")),
        "holdings": [
            {
                "name": row.get("company"),
                "isin": row.get("isin"),
                "country": row.get("country"),
                "sector": row.get("sector"),
                "rating": row.get("rating"),
                "currency": row.get("currency"),
                "coupon": row.get("coupon"),
                "yield": row.get("yield"),
                "maturity": str(row.get("maturity")),
            }
            for row in rows_dict
        ],
    }


def pct(x: float) -> float:
    # 6dp, not 4: a genuinely nonzero weight this small (e.g. a residual
    # CAD cash balance at ~0.00003%) must survive rounding as a small
    # positive number, not collapse to exactly 0.0 -- the build scripts'
    # display layer (fmt_pct) is what decides whether to show "<0.1%" for
    # it, and it can only do that if this function hasn't already erased
    # the distinction between "genuinely zero" and "just very small".
    return round(x * 100, 6)


def val_eur(row: dict) -> float:
    v = row.get("valuation + accr. interest (eur)")
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def weight(row: dict) -> float:
    w = row.get("weight (%)")
    try:
        return float(w or 0) / 100.0
    except (TypeError, ValueError):
        return 0.0


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("excel_path")
    ap.add_argument("--profile", required=True, choices=RISK_PROFILES)
    ap.add_argument("--client-name", default="[Client Name]")
    ap.add_argument("--valuation-date", default=None, help="e.g. '6 August 2026'; defaults to today")
    ap.add_argument("--bucket-overrides", default=None)
    ap.add_argument("--market-cap-overrides", default=None)
    ap.add_argument("--sector-overrides", default=None)
    ap.add_argument("-o", "--output", default="parsed.json")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.excel_path, data_only=True)
    _, port_rows = load_sheet_rows(pick_portfolio_sheet(wb))
    proposed_bond_selection = parse_proposed_bonds(wb) if "Fixed Income" in wb.sheetnames else None

    bucket_overrides = load_bucket_overrides(args.bucket_overrides)
    mcap_overrides = load_market_cap_overrides(args.market_cap_overrides)
    sector_overrides = load_sector_overrides(args.sector_overrides)

    included, uncalled = [], []
    for row in port_rows:
        section = row.get("_section")
        if section not in PORTFOLIO_SECTIONS:
            continue
        w = weight(row)
        v = val_eur(row)
        if w <= 0 or v == 0:
            desc = row.get("description") or ""
            if "COMMIT" in str(desc).upper():
                uncalled.append({
                    "isin": row.get("isin code"),
                    "description": desc,
                    "committed_amount": row.get("balance"),
                    "currency": row.get("currency"),
                })
            continue
        row["_asset_class"] = (
            "Commodities" if is_precious_metal_account(row.get("description"), row.get("currency"))
            else ASSET_CLASS_MAP[section])
        row["_weight"] = w
        row["_value_eur"] = v
        try:
            row["_value_qc"] = float(row.get("valuation + accr. interest (in qc)") or 0)
        except (TypeError, ValueError):
            row["_value_qc"] = 0.0
        included.append(row)

    total_value_eur = sum(r["_value_eur"] for r in included)

    # Base/reporting currency: fixed at EUR, matching the custodian file's
    # own "Valuation + accr. interest (EUR)" column and its own stated
    # portfolio total (assumptions.md §3). This is NOT the same thing as
    # "most-held currency by weight" (that's currency_exposure_pct below,
    # a property of the holdings, not of the reporting base) -- conflating
    # the two mislabels every EUR-denominated aggregate figure with
    # whatever currency happens to dominate the book.
    base_currency = "EUR"
    ccy_val = defaultdict(float)
    for r in included:
        ccy_val[str(r.get("currency") or "Other")] += r["_value_eur"]
    dominant_holding_currency = max(ccy_val.items(), key=lambda kv: kv[1])[0] if ccy_val else base_currency

    # --- §5 asset allocation donut ---
    class_weight = defaultdict(float)
    for r in included:
        class_weight[r["_asset_class"]] += r["_weight"]
    asset_allocation = {k: pct(v) for k, v in sorted(class_weight.items(), key=lambda kv: -kv[1])}

    # --- currency exposure donut ---
    ccy_weight = defaultdict(float)
    for r in included:
        ccy_weight[str(r.get("currency") or "Other")] += r["_weight"]
    currency_exposure = {k: pct(v) for k, v in sorted(ccy_weight.items(), key=lambda kv: -kv[1])}

    # --- §7 geography / sector portfolio-wide ---
    # The Portfolio sheet carries no Sector column (unlike the separate
    # 'Fixed Income' proposal tab) -- sector exposure is only computable if
    # the caller supplies --sector-overrides; otherwise it is reported as
    # unavailable rather than fabricated. This matches the reference deck's
    # own disclosed caveat: "these fields are not in the client file."
    geo_weight = defaultdict(float)
    sector_weight = defaultdict(float)
    sector_data_available = bool(sector_overrides)
    for r in included:
        geo_weight[region_of(r.get("geographical breakdown"))] += r["_weight"]
        if sector_data_available:
            isin = str(r.get("isin code") or "")
            sector_weight[sector_overrides.get(isin, "Other")] += r["_weight"]
    geographic_exposure = {k: pct(v) for k, v in sorted(geo_weight.items(), key=lambda kv: -kv[1])}
    sector_exposure = (
        {k: pct(v) for k, v in sorted(sector_weight.items(), key=lambda kv: -kv[1])}
        if sector_data_available else None
    )

    # --- sleeves: top holdings + vehicle breakdown per asset class ---
    sleeves = {}
    for asset_class in sorted(set(r["_asset_class"] for r in included)):
        rows_ac = [r for r in included if r["_asset_class"] == asset_class]
        ac_total_w = sum(r["_weight"] for r in rows_ac)
        top = sorted(rows_ac, key=lambda r: -r["_weight"])[:5]
        vehicle_weight = defaultdict(float)
        for r in rows_ac:
            v = vehicle_of(r)
            vehicle_weight[v] += r["_weight"] / ac_total_w if ac_total_w else 0
        durations = [float(r["duration"]) for r in rows_ac if isinstance(r.get("duration"), (int, float))]
        n_with_dur = len(durations)
        sleeves[asset_class] = {
            "total_weight_pct": pct(ac_total_w),
            "num_lines": len(rows_ac),
            "top_holdings": [
                {
                    "name": r.get("description"),
                    "isin": r.get("isin code"),
                    "value_eur": r["_value_eur"],
                    "value_qc": r["_value_qc"],
                    "currency": r.get("currency"),
                    "weight_pct": pct(r["_weight"]),
                    "duration": r.get("duration"),
                    "yield_pct": r.get("yield (%)"),
                    "price": r.get("price (qc)"),
                    # §1a: this row's "quote currency" is a metal, so its
                    # value_qc is an ounce count, not an amount of money
                    "metal_account": is_precious_metal_account(r.get("description"), r.get("currency")),
                }
                for r in top
            ],
            "vehicle_breakdown_pct": {k: pct(v) for k, v in sorted(vehicle_weight.items(), key=lambda kv: -kv[1])},
            "duration_coverage": f"{n_with_dur} of {len(rows_ac)}" if asset_class == "Fixed Income" else None,
        }

    # --- §8 equity breakdown (renormalized to the equity sleeve) ---
    equity_rows = [r for r in included if r["_asset_class"] == "Equities"]
    eq_total_w = sum(r["_weight"] for r in equity_rows) or 1.0
    eq_geo, eq_sector, eq_mcap = defaultdict(float), defaultdict(float), defaultdict(float)
    mcap_available = bool(mcap_overrides)
    for r in equity_rows:
        w_norm = r["_weight"] / eq_total_w
        eq_geo[equity_region_of(r.get("geographical breakdown"))] += w_norm
        if sector_data_available:
            isin = str(r.get("isin code") or "")
            eq_sector[sector_overrides.get(isin, "Other")] += w_norm
        if mcap_available:
            isin = str(r.get("isin code") or "")
            eq_mcap[mcap_overrides.get(isin, "Large cap")] += w_norm
    equity_breakdown = {
        "by_geography_pct": {k: pct(v) for k, v in sorted(eq_geo.items(), key=lambda kv: -kv[1])},
        "by_sector_pct": ({k: pct(v) for k, v in sorted(eq_sector.items(), key=lambda kv: -kv[1])} if sector_data_available else None),
        "by_market_cap_pct": ({k: pct(v) for k, v in sorted(eq_mcap.items(), key=lambda kv: -kv[1])} if mcap_available else None),
    }

    # --- §9 concentration & top holdings ---
    by_weight = sorted(included, key=lambda r: -r["_weight"])
    cum = 0.0
    concentration_rows = []
    for i, r in enumerate(by_weight[:20], 1):
        cum += r["_weight"]
        concentration_rows.append({
            "rank": i, "name": r.get("description"), "value_eur": r["_value_eur"],
            "weight_pct": pct(r["_weight"]), "cumulative_pct": pct(cum),
        })
    concentration = {
        "top_holdings": concentration_rows,
        "top_5_pct": pct(sum(r["_weight"] for r in by_weight[:5])),
        "top_10_pct": pct(sum(r["_weight"] for r in by_weight[:10])),
        "top_20_pct": pct(sum(r["_weight"] for r in by_weight[:20])),
        "positions_above_5pct": sum(1 for r in included if r["_weight"] > 0.05),
    }

    # --- §10 liquidity profile ---
    liq_weight = defaultdict(float)
    for r in included:
        liq_weight[classify_liquidity(r, r["_asset_class"], vehicle_of(r))] += r["_weight"]
    liquidity_profile = {k: pct(v) for k, v in sorted(liq_weight.items(), key=lambda kv: -kv[1])}

    # --- §11 income & rate sensitivity ---
    fi_rows = [r for r in included if r["_asset_class"] == "Fixed Income"]
    running_yield_num, running_yield_den = 0.0, 0.0
    for r in included:
        y = r.get("yield (%)")
        if isinstance(y, (int, float)):
            running_yield_num += r["_weight"] * y
            running_yield_den += r["_weight"]
    running_yield_pct = round(running_yield_num, 4) if running_yield_den else None

    fi_total_val = sum(r["_value_eur"] for r in fi_rows)
    dur_num, dur_den = 0.0, 0.0
    duration_subsleeves = defaultdict(lambda: {"value": 0.0, "dur_weighted": 0.0})
    for r in fi_rows:
        d = r.get("duration")
        if isinstance(d, (int, float)):
            dur_num += r["_value_eur"] * d
            dur_den += r["_value_eur"]
        label = describe_bond_subsleeve(r)
        duration_subsleeves[label]["value"] += r["_value_eur"]
        if isinstance(d, (int, float)):
            duration_subsleeves[label]["dur_weighted"] += r["_value_eur"] * d
    fi_duration = round(dur_num / dur_den, 2) if dur_den else None
    rate_impact_eur = round(-fi_duration * 0.01 * fi_total_val, 2) if fi_duration else None

    income_by_sleeve = defaultdict(float)
    for r in included:
        y = r.get("yield (%)")
        c = r.get("coupon (%)")
        rate = y if isinstance(y, (int, float)) else (c if isinstance(c, (int, float)) else None)
        if rate is not None:
            income_by_sleeve[r["_asset_class"]] += r["_value_eur"] * rate / 100.0
    total_income = sum(income_by_sleeve.values()) or 1.0

    income = {
        "running_yield_pct": running_yield_pct,
        "fi_duration_years": fi_duration,
        "rate_impact_100bp_eur": rate_impact_eur,
        "income_by_sleeve": [
            {"sleeve": k, "income_eur": round(v, 2), "pct_of_income": pct(v / total_income)}
            for k, v in sorted(income_by_sleeve.items(), key=lambda kv: -kv[1])
        ],
        "fi_duration_subsleeves": [
            {
                "label": k,
                "value_eur": round(v["value"], 2),
                "avg_duration": round(v["dur_weighted"] / v["value"], 2) if v["value"] and v["dur_weighted"] else 0.0,
                "contribution": round(v["dur_weighted"] / fi_total_val, 2) if fi_total_val else 0.0,
            }
            for k, v in sorted(duration_subsleeves.items(), key=lambda kv: -kv[1]["value"])
        ],
    }

    # --- §12 risk profile / growth-defensive split ---
    growth_w = sum(v for k, v in class_weight.items() if k in GROWTH_CLASSES)
    defensive_w = sum(v for k, v in class_weight.items() if k in DEFENSIVE_CLASSES)
    risk_profile = {
        "selected_profile": args.profile,
        "growth_assets_pct": pct(growth_w),
        "defensive_assets_pct": pct(defensive_w),
    }

    # --- line items table source (assumptions.md §4, slides 9-10) ---
    line_items = []
    fi_bucket_order = ["Govies 1-10 (local)", "Govies 10+ (local)", "High Yield (local or global hdg)",
                        "Corporate IG (local)", "EM Debt"]
    equity_region_order = ["US", "Eurozone", "UK", "Switzerland", "Japan", "EM", "Global / thematic"]
    for r in included:
        ac = r["_asset_class"]
        if ac == "Fixed Income":
            group = classify_fi_bucket(r, bucket_overrides)
        elif ac == "Cash":
            group = "Cash"
        elif ac == "Equities":
            group = equity_region_of(r.get("geographical breakdown"))
        else:
            group = ac
        line_items.append({
            "top_category": ac,
            "group": group,
            "isin": r.get("isin code"),
            "instrument": r.get("description"),
            "weight_pct": pct(r["_weight"]),
        })
    # stable sort: top_category (Cash, Fixed Income, Equities, Alternatives,
    # Structured Products, Private Assets, Commodities, Other), then FI bucket
    # order, then descending weight within group.
    top_cat_order = ["Cash", "Fixed Income", "Equities", "Alternatives", "Structured Products",
                      "Private Assets", "Commodities", "Other"]

    def sort_key(li):
        tc = top_cat_order.index(li["top_category"]) if li["top_category"] in top_cat_order else 99
        if li["top_category"] == "Fixed Income":
            gi = fi_bucket_order.index(li["group"]) if li["group"] in fi_bucket_order else 99
        elif li["top_category"] == "Equities":
            gi = equity_region_order.index(li["group"]) if li["group"] in equity_region_order else 99
        else:
            gi = 0
        return (tc, gi, -li["weight_pct"])

    line_items.sort(key=sort_key)

    import datetime as _dt
    valuation_date = args.valuation_date or _dt.date.today().strftime("%-d %B %Y")

    parsed = {
        "client_name": args.client_name,
        "risk_profile_input": args.profile,
        "valuation_date": valuation_date,
        "base_currency": base_currency,
        "dominant_holding_currency": dominant_holding_currency,
        "total_value_eur": round(total_value_eur, 2),
        "num_positions": len(included),
        "largest_position_pct": pct(by_weight[0]["_weight"]) if by_weight else 0,
        "liquid_share_pct": pct(sum(v for k, v in liq_weight.items() if k != ILLIQUID_BUCKET)),
        "asset_allocation_pct": asset_allocation,
        "currency_exposure_pct": currency_exposure,
        "geographic_exposure_pct": geographic_exposure,
        "sector_exposure_pct": sector_exposure,
        "sleeves": sleeves,
        "equity_breakdown": equity_breakdown,
        "concentration": concentration,
        "liquidity_profile_pct": liquidity_profile,
        "income": income,
        "risk_profile": risk_profile,
        "uncalled_commitments": uncalled,
        "line_items": line_items,
        "proposed_bond_selection": proposed_bond_selection,
    }

    Path(args.output).write_text(json.dumps(parsed, indent=2, default=str))
    print(f"Wrote {args.output}: {len(included)} positions, "
          f"{parsed['total_value_eur']:,.0f} EUR, dominant holding currency {dominant_holding_currency}")


if __name__ == "__main__":
    main()
