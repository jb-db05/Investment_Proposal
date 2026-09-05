#!/usr/bin/env python3
"""
Main orchestrator: portfolio Excel + market PDF summary + risk profile ->
a fully populated Investment Proposal deck.

    python build_proposal.py \
        --excel Portfolio.xlsx --profile Balanced \
        --client-name "Jane Doe" \
        --market-update market_update.json \
        -o "Investment Proposal - Jane Doe.pptx"

What this script does NOT do: read the market PDF itself. Extracting an
arbitrary week's commentary into structured slide content is a
summarization/judgment task, not a deterministic transform — see
SKILL.md and references/slide_recipe.md for how Claude produces
market_update.json from the PDF before this script runs. Everything
downstream of that JSON (and of parsed.json from parse_portfolio.py) is
a deterministic, no-invented-numbers fill of the template.

Text formatting rule (important): every fill_* function below sets text
via set_text()/set_shape_lines(), which preserve each shape's own
per-paragraph run formatting (size, bold, color, font) by position — never
python-pptx's `text_frame.text = ...`, which collapses everything to one
unstyled run. The template's own designer already put the right formatting
on each line (e.g. a stat box's number is bold orange, its caption is
navy 12pt) — these helpers just swap the text, never the look, unless a
`bold_overrides` dict explicitly forces a specific line's boldness.
"""
from __future__ import annotations

import argparse
import copy as _copy
import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
import build_line_items_tables as line_items_mod  # noqa: E402
import update_chart as chart_mod  # noqa: E402

TEMPLATE_PATH = SCRIPT_DIR.parent / "assets" / "template.pptx"

# Brand palette, pulled from the template's own theme (ppt/theme/theme1.xml,
# "Syz 2024") and from the exact colors already used on its shapes — never
# invented. See references/assumptions.md's closing section for how each
# value was found.
NAVY = RGBColor(0x20, 0x29, 0x45)        # dk2 — the deck's real "black": use this, never 000000
GOLD = RGBColor(0xFF, 0xC5, 0x45)        # lt2 — table header fill, Cash category bar
SLATE = RGBColor(0x4B, 0x5F, 0x80)       # accent6 — Fixed Income / Equity category bars
MINT = RGBColor(0x3B, 0xAF, 0x90)        # accent2 — Commodities category bar; "positive" scoreboard color
ORANGE = RGBColor(0xFF, 0xA4, 0x00)      # accent3 — profile-dial highlight; "Other" category bar
TIGER = RGBColor(0xFF, 0x6C, 0x0E)       # accent4 — "negative" scoreboard color; KPI-box numbers
MAUVE = RGBColor(0xAC, 0x5D, 0x85)       # accent5 — Alternatives / Private Assets category bar
SKY = RGBColor(0x79, 0xD6, 0xFF)         # accent1 — Structured Products category bar
DOT_DEFAULT = RGBColor(0x3A, 0xB7, 0xC8)  # literal color already used for the non-selected profile dots

RISK_PROFILE_CHARACTERISTICS = {
    "Fixed Income":  {"horizon": "Short term (1 to 3 years)",       "liquidity": "High liquidity, capital preservation priority"},
    "Conservative":  {"horizon": "Short to medium term (3 to 5 years)", "liquidity": "Regular distribution, stable capital"},
    "Moderate":      {"horizon": "Medium to long term (5 to 10 years)", "liquidity": "Regular distribution but stable capital"},
    "Balanced":      {"horizon": "Medium to long term (5 to 10 years)", "liquidity": "Balanced between growth and distribution"},
    "Growth":        {"horizon": "Long term (7 to 12 years)",       "liquidity": "Growth priority, distributions opportunistic"},
    "Equity":        {"horizon": "Long term (10+ years)",           "liquidity": "Full growth priority, minimal liquidity needs"},
}
# assumptions.md §12 / slide8_profile_dial.md: growth-asset band each profile is centered on.
RISK_PROFILE_GROWTH_BAND = {
    "Fixed Income": (0, 10), "Conservative": (10, 25), "Moderate": (25, 45),
    "Balanced": (45, 65), "Growth": (65, 85), "Equity": (85, 100),
}
# Slide 8's risk-return graph: one oval per profile, inside group "Group 3".
# Shape names extracted directly from the template — see references/assumptions.md.
PROFILE_DOT_SHAPE = {
    "Fixed Income": "Oval 8",
    "Conservative": "Oval 9",
    "Moderate": "Oval 11",
    "Balanced": "Oval 17",
    "Growth": "Oval 13",
    "Equity": "Oval 15",
}
# The profile-name captions next to each dot. Only 4 of the 6 had explicit
# styling in the template (teal, bold, 11pt) — Moderate and Balanced had no
# <a:rPr> at all, inheriting an unstyled default. All six now get explicit
# styling every build, keyed off the selected profile, rather than leaving
# two of them to whatever the template happens to default to.
PROFILE_LABEL_SHAPE = {
    "Fixed Income": "TextBox 7",
    "Conservative": "TextBox 10",
    "Moderate": "TextBox 12",
    "Balanced": "TextBox 18",
    "Growth": "TextBox 14",
    "Equity": "TextBox 16",
}
# Sleeve slides where the "largest holdings" table's Yield/Coupon column is
# always empty (equities, alternatives and commodities don't carry
# yield/coupon data) -- dropped per user instruction rather than shown blank.
SLEEVES_WITHOUT_YIELD_COLUMN = {"Equities", "Private Assets", "Commodities"}


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_shape_lines(shape, lines, bold_overrides=None):
    """Set a shape's text to `lines` (one paragraph per line), preserving
    each existing paragraph's own run formatting (size/color/bold/font) by
    position — only the text changes. If there are more lines than existing
    paragraphs, extra paragraphs are cloned from the last existing one; if
    fewer, trailing paragraphs are removed. `bold_overrides` is an optional
    {line_index: bool} dict to force specific lines bold/not-bold regardless
    of what the template had (e.g. slide 4's driver boxes: title stays
    however the template had it, every other line is forced non-bold)."""
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    lines = [str(l) for l in lines]
    bold_overrides = bold_overrides or {}

    existing = list(tf.paragraphs)
    n_existing, n_lines = len(existing), len(lines)
    if n_lines > n_existing and n_existing > 0:
        anchor = existing[-1]._p
        for _ in range(n_lines - n_existing):
            clone = _copy.deepcopy(anchor)
            anchor.addnext(clone)
            anchor = clone
    elif n_lines < n_existing:
        for extra in list(tf.paragraphs)[n_lines:]:
            extra._p.getparent().remove(extra._p)

    for i, (para, line) in enumerate(zip(tf.paragraphs, lines)):
        if not para.runs:
            para.add_run()
        for run in list(para.runs)[1:]:
            run._r.getparent().remove(run._r)
        para.runs[0].text = line
        if i in bold_overrides:
            para.runs[0].font.bold = bold_overrides[i]


def set_text(slide, name, text, bold_overrides=None):
    """set_shape_lines() by shape name, splitting text on '\\n'."""
    set_shape_lines(find_shape(slide, name), str(text).split("\n"), bold_overrides)


# The subtitle placeholder ("Text Placeholder 2", just under every slide
# title) uses <a:spAutoFit/> with anchor="b" (bottom) at the layout level:
# the shape's own height grows to fit its text, and because it's anchored
# to the bottom of a box whose top edge is fixed, a subtitle that wraps to
# a second line visibly shifts *down* rather than growing upward — exactly
# the bug reported. Editing text in place can't fix that; the shape itself
# has to go. SUBTITLE_LEFT/TOP match the template's own position (H=1.85cm,
# V=4.8cm) exactly, but as a plain textbox with a fixed height, top-anchored
# text and no autofit, so it never moves regardless of line count.
SUBTITLE_LEFT = Emu(665163)
SUBTITLE_TOP = Emu(1728947)
SUBTITLE_WIDTH = Emu(9398001)
SUBTITLE_HEIGHT = Emu(560000)
SUBTITLE_NAME = "Text Placeholder 2"


def set_subtitle(slide, text):
    old = find_shape(slide, SUBTITLE_NAME)
    if old is not None:
        old._element.getparent().remove(old._element)
    tb = slide.shapes.add_textbox(SUBTITLE_LEFT, SUBTITLE_TOP, SUBTITLE_WIDTH, SUBTITLE_HEIGHT)
    tb.name = SUBTITLE_NAME
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    return tb


def color_scoreboard_column(shape, values):
    """Color each value paragraph (paragraph index 1+, index 0 is the
    column header) green if positive, tiger-orange if negative — matching
    the sign of the value string itself, not a guess."""
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    for i, val in enumerate(values, start=1):
        if i >= len(tf.paragraphs) or not tf.paragraphs[i].runs:
            continue
        run = tf.paragraphs[i].runs[0]
        v = str(val).strip()
        if v.startswith("-"):
            run.font.color.rgb = TIGER
        elif v.startswith("+"):
            run.font.color.rgb = MINT


# Rough average glyph advance for the deck's body font, as a fraction of the
# font size. Only used to decide whether a scoreboard value fits its column —
# an approximation, deliberately generous, never a layout measurement.
AVG_CHAR_EM = 0.52
COLUMN_PADDING_PT = 14   # a little slack, since the advance above is an estimate


def _needed_width(values, size_pt):
    return Pt(max(len(str(v)) for v in values) * size_pt * AVG_CHAR_EM + COLUMN_PADDING_PT)


def fit_scoreboard_columns(label_box, boxes, columns, size_pt=11):
    """Re-width slide 3's three scoreboard columns to their own longest value.

    The template sizes them for an index name and two '+13.70%' figures, so a
    source whose values are words ('Underweight') overflows the two narrow
    columns. Columns are re-laid left to right from the first column's left
    edge, keeping the template's own gaps and staying inside the width of the
    block's label ('Cross-asset scoreboard'), which is what defines the
    block's right edge. No-op when the template's widths already fit — the
    weekly-scoreboard case, which keeps the reference deck's exact layout."""
    needed = [int(_needed_width(col, size_pt)) for col in columns]
    # A column only ever grows past the template's own width; one that already
    # has room keeps it, so the weekly-scoreboard layout is untouched.
    targets = [max(n, b.width) for n, b in zip(needed, boxes)]
    if targets == [b.width for b in boxes]:
        return
    gaps = [boxes[1].left - (boxes[0].left + boxes[0].width),
            boxes[2].left - (boxes[1].left + boxes[1].width)]
    available = (label_box.left + label_box.width) - boxes[0].left - sum(gaps)
    # Growing one column past the block's right edge is paid for by columns
    # that are wider than their own content needs — the label column, usually,
    # which the template sizes for long index names.
    over = sum(targets) - available
    for i in sorted(range(len(targets)), key=lambda i: needed[i] - targets[i]):
        if over <= 0:
            break
        take = min(targets[i] - needed[i], over)
        targets[i] -= take
        over -= take
    if over > 0:  # every column is at its minimum: shrink them all in proportion
        targets = [int(t * available / sum(targets)) for t in targets]
    left = boxes[0].left
    for box, width, gap in zip(boxes, targets, gaps + [0]):
        box.left, box.width = Emu(int(left)), Emu(int(width))
        left += int(width) + gap


def recolor_column(shape, color):
    """Force every value paragraph (index 1+, index 0 is the column header)
    to one color — the counterpart of color_scoreboard_column() for columns
    whose values aren't signed numbers."""
    if shape is None or not shape.has_text_frame:
        return
    for para in list(shape.text_frame.paragraphs)[1:]:
        for run in para.runs:
            run.font.color.rgb = color


DOT_SIZE_DEFAULT = Emu(250065)     # matches every non-selected dot in the template
DOT_SIZE_HIGHLIGHT = Emu(340000)   # a bit bigger than the template's own highlight size (314873)


def _resize_dot_keep_center(shape, new_size: Emu):
    cx = shape.left + shape.width / 2
    cy = shape.top + shape.height / 2
    shape.width = new_size
    shape.height = new_size
    shape.left = Emu(int(cx - new_size / 2))
    shape.top = Emu(int(cy - new_size / 2))


def recolor_profile_dial(slide, selected_profile):
    """Slide 8's risk-return graph: the selected profile's dot AND its name
    caption turn orange and a bit larger (the template's own highlight
    treatment); every other profile's dot and caption -- including
    Moderate when it's not selected -- reverts to the default teal at the
    default size."""
    group = find_shape(slide, "Group 3")
    if group is None:
        return
    for sub in group.shapes:
        if sub.name.startswith("Oval"):
            sub.fill.solid()
            sub.fill.fore_color.rgb = DOT_DEFAULT
            _resize_dot_keep_center(sub, DOT_SIZE_DEFAULT)
    for profile, label_name in PROFILE_LABEL_SHAPE.items():
        # these TextBoxes live nested inside "Group 3", same as the ovals --
        # find_shape(slide, ...) only searches top-level shapes and would
        # silently miss them
        label = next((sub for sub in group.shapes if sub.name == label_name), None)
        if label is None or not label.has_text_frame or not label.text_frame.paragraphs[0].runs:
            continue
        run = label.text_frame.paragraphs[0].runs[0]
        selected = profile == selected_profile
        run.font.color.rgb = ORANGE if selected else DOT_DEFAULT
        run.font.size = Pt(12) if selected else Pt(11)
        run.font.bold = selected
    highlight_name = PROFILE_DOT_SHAPE.get(selected_profile)
    if highlight_name:
        dot = next((sub for sub in group.shapes if sub.name == highlight_name), None)
        if dot is not None:
            dot.fill.solid()
            dot.fill.fore_color.rgb = ORANGE
            _resize_dot_keep_center(dot, DOT_SIZE_HIGHLIGHT)


def fmt_pct(value: float) -> str:
    """1-decimal percent, except a genuinely nonzero value that rounds to
    0.0% shows as '<0.1%' instead -- '0.0%' next to a real position (e.g.
    "INCOME TO BE RECEIVED IN GBP  0.0%") reads as an error or missing
    data, not as a very small real number."""
    if 0 < value < 0.1:
        return "<0.1%"
    return f"{value:.1f}%"


def fmt_money_m(value_eur: float, currency: str) -> str:
    return f"{currency} {value_eur / 1_000_000:.1f}m"


def fmt_money_k(value: float, currency: str) -> str:
    return f"{currency} {value / 1_000:.0f}k"


def get_slide(prs, index_1based):
    return list(prs.slides)[index_1based - 1]


# ---------------------------------------------------------------- charts ---

def update_donut_by_name(slide, chart_name, data, label_style="whole"):
    sh = find_shape(slide, chart_name)
    if sh is None or not sh.has_chart:
        return
    chart_mod.update_donut(sh, data, label_style=label_style)


def update_bar_by_name(slide, chart_name, data):
    sh = find_shape(slide, chart_name)
    if sh is None or not sh.has_chart:
        return
    chart_mod.update_bar(sh, data)


# ----------------------------------------------------------------- tables --
# Table header style matches the template's own tables exactly (all of
# slides 13/15/17/18/19/22 use gold fill + navy bold text for headers —
# see references/assumptions.md).
CARD_FILL = RGBColor(0xF2, 0xF4, 0xF7)   # the pale panel behind the sleeve slides' text blocks
TABLE_HEADER_FILL = GOLD
TABLE_HEADER_FONT = NAVY


def _cell(cell, text, *, bold=False, fill=None, font_color=NAVY, size=10, align=PP_ALIGN.LEFT):
    cell.margin_left = Emu(45720)
    cell.margin_right = Emu(45720)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if not p.runs:
        p.add_run()
    for run in list(p.runs)[1:]:
        run._r.getparent().remove(run._r)
    p.runs[0].text = str(text)
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.bold = bold
    p.runs[0].font.color.rgb = font_color


def rebuild_holdings_table(slide, table_shape_name, holdings, currency_fallback="USD", include_yield_column=True):
    """Replace the sleeve's 'largest holdings' table with one sized to the
    actual number of top holdings (assumptions.md doesn't fabricate empty
    rows, and the reference template's tables are too small for a sleeve
    with more than 1-5 real holdings). The Yield/Coupon column is only
    included where the sleeve can actually carry that data."""
    sh = find_shape(slide, table_shape_name)
    if sh is None or not sh.has_table:
        return
    left, top, width, height = sh.left, sh.top, sh.width, sh.height
    sh._element.getparent().remove(sh._element)

    n_cols = 4 if include_yield_column else 3
    n_rows = 1 + max(1, len(holdings))
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = graphic_frame.table
    col_fracs = (0.46, 0.18, 0.14, 0.22) if include_yield_column else (0.52, 0.24, 0.24)
    for i, f in enumerate(col_fracs):
        table.columns[i].width = Emu(int(width * f))
    headers = ["Holding", "Value", "Share"] + (["Yield / Coupon"] if include_yield_column else [])
    for c, h in enumerate(headers):
        _cell(table.cell(0, c), h, bold=True, fill=TABLE_HEADER_FILL, font_color=TABLE_HEADER_FONT, size=9.5)
    if not holdings:
        _cell(table.cell(1, 0), "No holdings in this sleeve", size=9)
        for c in range(1, n_cols):
            _cell(table.cell(1, c), "")
        return
    for r, h in enumerate(holdings, 1):
        # A precious-metal account quotes in ounces (assumptions.md §1a), so
        # its "value in quote currency" is an ounce count, not an amount of
        # money — 144 XPT would print as "XPT 0k" for a 220k position. Those
        # rows are shown in the portfolio's base currency instead.
        if h.get("metal_account"):
            ccy, value = currency_fallback, h.get("value_eur") or 0
        else:
            ccy = h.get("currency") or currency_fallback
            value = h.get("value_qc") or h.get("value_eur") or 0
        _cell(table.cell(r, 0), h.get("name") or "", size=9)
        _cell(table.cell(r, 1), fmt_money_k(value, ccy), size=9)
        _cell(table.cell(r, 2), fmt_pct(h['weight_pct']), size=9, align=PP_ALIGN.RIGHT)
        if include_yield_column:
            yld = h.get("yield_pct")
            _cell(table.cell(r, 3), (f"{yld:.2f}%" if isinstance(yld, (int, float)) else "—"), size=9, align=PP_ALIGN.RIGHT)


def fill_concentration_table(slide, table_shape_name, rows, currency="EUR"):
    sh = find_shape(slide, table_shape_name)
    if sh is None or not sh.has_table:
        return
    table = sh.table
    max_rows = len(table.rows) - 1
    for i in range(max_rows):
        r = i + 1
        if i < len(rows):
            row = rows[i]
            _cell(table.cell(r, 0), str(row["rank"]), size=9)
            _cell(table.cell(r, 1), row["name"] or "", size=9)
            _cell(table.cell(r, 2), f"{row['value_eur']:,.0f}", size=9, align=PP_ALIGN.RIGHT)
            _cell(table.cell(r, 3), fmt_pct(row['weight_pct']), size=9, align=PP_ALIGN.RIGHT)
            _cell(table.cell(r, 4), f"{row['cumulative_pct']:.1f}%", size=9, align=PP_ALIGN.RIGHT)
        else:
            for c in range(5):
                _cell(table.cell(r, c), "")


def fill_income_tables(slide, sleeve_table_name, duration_table_name, income):
    income_row_height = None
    sh = find_shape(slide, sleeve_table_name)
    if sh is not None and sh.has_table:
        table = sh.table
        income_row_height = table.rows[0].height
        rows = income["income_by_sleeve"]
        for i in range(len(table.rows) - 1):
            r = i + 1
            if i < len(rows):
                row = rows[i]
                _cell(table.cell(r, 0), row["sleeve"], size=9)
                _cell(table.cell(r, 1), f"{row['income_eur']:,.0f}", size=9, align=PP_ALIGN.RIGHT)
                _cell(table.cell(r, 2), fmt_pct(row['pct_of_income']), size=9, align=PP_ALIGN.RIGHT)
            else:
                for c in range(3):
                    _cell(table.cell(r, c), "")

    # Rebuilt with the exact same header style (gold fill / navy text) AND
    # the exact same width/row-height as the income-by-sleeve table right
    # next to it -- these two tables must look like one system, not two
    # different designs at two different scales.
    sh2 = find_shape(slide, duration_table_name)
    if sh2 is not None and sh2.has_table:
        left, top, height = sh2.left, sh2.top, sh2.height
        width = sh.width if sh is not None else sh2.width
        sh2._element.getparent().remove(sh2._element)
        rows = income["fi_duration_subsleeves"]
        n_rows = 1 + max(1, len(rows))
        gf = slide.shapes.add_table(n_rows, 4, left, top, width, height)
        table = gf.table
        if income_row_height:
            for r in table.rows:
                r.height = income_row_height
        headers = ["Sub-sleeve", "Value", "Avg duration", "Contribution"]
        for c, h in enumerate(headers):
            _cell(table.cell(0, c), h, bold=True, fill=TABLE_HEADER_FILL, font_color=TABLE_HEADER_FONT, size=9)
        for r, row in enumerate(rows, 1):
            _cell(table.cell(r, 0), row["label"], size=8.5)
            _cell(table.cell(r, 1), f"{row['value_eur']:,.0f}", size=8.5, align=PP_ALIGN.RIGHT)
            _cell(table.cell(r, 2), f"{row['avg_duration']:.1f}", size=8.5, align=PP_ALIGN.RIGHT)
            _cell(table.cell(r, 3), f"{row['contribution']:.2f}", size=8.5, align=PP_ALIGN.RIGHT)


# ------------------------------------------------------------- slide fills -

def fill_cover(prs, parsed):
    slide = get_slide(prs, 1)
    ccy = parsed["base_currency"]
    amount_m = parsed["total_value_eur"] / 1_000_000
    set_text(slide, "Text Placeholder 5",
             f"{parsed['risk_profile_input']} {ccy} {amount_m:.1f}m\n\n"
             f"Prepared for {parsed['client_name']}\n\n"
             f"Your Relationship Manager: [Name]\n\n"
             f"Your Advisor: ")


def fill_profile_dial(prs, parsed):
    slide = get_slide(prs, 8)
    profile = parsed["risk_profile_input"]
    growth_pct = parsed["risk_profile"]["growth_assets_pct"]
    set_text(slide, "TextBox 23", f"Proposed profile for this portfolio: {profile} - {growth_pct:.0f}% in growth assets")
    chars = RISK_PROFILE_CHARACTERISTICS[profile]
    sh = find_shape(slide, "Table 44")
    if sh is not None and sh.has_table:
        table = sh.table
        mapping = {
            "Investment horizon": chars["horizon"],
            "Reference of currency": parsed["base_currency"],
            "Risk level": profile,
            "Liquidity needs": chars["liquidity"],
        }
        for row in table.rows:
            label = row.cells[0].text.strip()
            if label in mapping:
                _cell(row.cells[1], mapping[label], size=10)
    recolor_profile_dial(slide, profile)


def fill_portfolio_overview(prs, parsed):
    slide = get_slide(prs, 11)
    ccy = parsed["base_currency"]
    set_text(slide, "Rectangle 3", f"{ccy} {parsed['total_value_eur']/1_000_000:.1f}M\nTotal value")
    set_text(slide, "Rectangle 4", f"{parsed['num_positions']}\nPositions")
    set_text(slide, "Rectangle 5", f"{parsed['largest_position_pct']:.1f}%\nLargest position")
    set_text(slide, "Rectangle 6", f"{parsed['liquid_share_pct']:.1f}%\nLiquid share")
    ry = parsed["income"]["running_yield_pct"]
    set_text(slide, "Rectangle 7", f"{ry:.1f}%\nRunning yield" if ry is not None else "n/a\nRunning yield")
    update_donut_by_name(slide, "Chart 10", parsed["asset_allocation_pct"])
    update_donut_by_name(slide, "Chart 14", parsed["currency_exposure_pct"])


def set_donut_legend(slide, textbox_names, data: dict[str, float]):
    """Fill the fixed set of companion 'legend' textboxes next to a donut
    with one-decimal percent labels, in data order. If there are more/fewer
    categories than textboxes (common — the template was sized for one
    example's category count), fill what fits and blank the rest rather
    than leaving stale numbers from the reference client behind."""
    items = list(data.items())
    for i, name in enumerate(textbox_names):
        text = chart_mod.pct_label_one_decimal(*items[i]) if i < len(items) else ""
        set_text(slide, name, text)


GEOGRAPHIC_SLIDE = 12
EQUITY_BREAKDOWN_SLIDE = 16


def fill_geographic_exposure(prs, parsed):
    """Both geographic views on one slide: the whole portfolio on the left,
    the equity sleeve on the right.

    The template's right-hand panel was "Sector exposure", and the equity
    sleeve's regional split had a slide of its own. Sector and market-cap
    breakdowns are out of this deck entirely (see build(): the equity-
    breakdown slide is dropped), so the freed panel carries the equity
    geography instead of a chart that could only ever say "Not available"
    without a desk sector mapping — assumptions.md §7."""
    slide = get_slide(prs, GEOGRAPHIC_SLIDE)
    sleeve = parsed["sleeves"].get("Equities", {})
    ccy = parsed["base_currency"]
    val_m = (sleeve.get("total_weight_pct", 0) / 100) * parsed["total_value_eur"] / 1_000_000
    set_text(slide, "Title 1", "Geographic exposure")
    set_subtitle(slide,
                 f"Regional breakdown of the whole portfolio and, separately, of the equity "
                 f"sleeve — {sleeve.get('num_lines', 0)} lines, {ccy} {val_m:.1f}m, "
                 f"{sleeve.get('total_weight_pct', 0):.1f}% of the portfolio.")
    set_text(slide, "TextBox 3", "Whole portfolio")
    update_bar_by_name(slide, "Chart 5", parsed["geographic_exposure_pct"])
    set_text(slide, "TextBox 6", "Equity sleeve")
    update_bar_by_name(slide, "Chart 8", parsed["equity_breakdown"]["by_geography_pct"])


SLEEVE_SLIDES = {
    # slide_number: parsed.json sleeve key
    13: "Fixed Income",
    15: "Equities",
    17: "Private Assets",
    18: "Commodities",
    19: "Structured Products",
}


def fill_sleeve_slides(prs, parsed):
    for slide_num, sleeve_key in SLEEVE_SLIDES.items():
        slide = get_slide(prs, slide_num)
        sleeve = parsed["sleeves"].get(sleeve_key)
        if sleeve is None:
            set_text(slide, "TextBox 5", "No holdings in this sleeve for this portfolio.")
            continue
        include_yield = sleeve_key not in SLEEVES_WITHOUT_YIELD_COLUMN
        rebuild_holdings_table(slide, "Table 6", sleeve["top_holdings"], parsed["base_currency"],
                                include_yield_column=include_yield)
        vehicle_chart = find_shape(slide, "Chart 9")
        if vehicle_chart is not None and vehicle_chart.has_chart:
            update_donut_by_name(slide, "Chart 9", sleeve["vehicle_breakdown_pct"])
        cov = sleeve.get("duration_coverage")
        caveat_box = find_shape(slide, "TextBox 11")
        if caveat_box is not None:
            if sleeve_key == "Fixed Income" and cov:
                set_text(slide, "TextBox 11",
                          f"⚠ Duration coverage: {cov} fixed-income line(s) report a modified duration; "
                          f"portfolio duration reflects only those.")
            else:
                set_text(slide, "TextBox 11", "")


# Original template position of the "Fixed Income Breakdown" (proposed bond
# selection) slide. It is dropped when the Excel carries no 'Fixed Income'
# proposal tab — leaving it would ship the template's own example selection,
# six issues and EUR 600,000 of bonds, as if proposed for this client. See
# the removal list in build().
PROPOSED_BOND_SLIDE = 14


def fill_proposed_bond_selection(prs, parsed):
    pb = parsed.get("proposed_bond_selection")
    if not pb:
        return  # slide removed later, see build()
    slide = get_slide(prs, PROPOSED_BOND_SLIDE)
    set_text(slide, "Text 1",
             f"Proposed bond selection: {pb['num_issues']} issues, "
             f"{pb['currency']} {pb['total_amount']:,.0f}, equally weighted at "
             f"{pb['currency']} {pb['amount_each']:,.0f} each.")
    update_donut_by_name(slide, "Chart 0", pb["by_geography_pct"], label_style="one_decimal")
    update_donut_by_name(slide, "Chart 1", pb["by_currency_pct"], label_style="one_decimal")
    update_donut_by_name(slide, "Chart 2", pb["by_maturity_pct"], label_style="one_decimal")
    update_donut_by_name(slide, "Chart 3", pb["by_rating_pct"], label_style="one_decimal")
    set_text(slide, "Text 16", "By currency")
    set_donut_legend(slide, ["Text 5", "Text 7", "Text 9", "Text 11", "Text 13", "Text 15"], pb["by_geography_pct"])
    set_donut_legend(slide, ["Text 18", "Text 20", "Text 22"], pb["by_currency_pct"])
    set_donut_legend(slide, ["Text 25", "Text 27", "Text 29"], pb["by_maturity_pct"])
    set_donut_legend(slide, ["Text 32", "Text 34", "Text 36", "Text 38", "Text 40"], pb["by_rating_pct"])
    # internal data-sourcing footnote — not for the client's eyes
    set_text(slide, "Text 41", "")
    # ...and the source deck's hardcoded "12" in the bottom-right corner: a
    # page number left over from its own pagination, already wrong before this
    # skill existed and carried by no other data slide. Blanked, not
    # renumbered.
    set_text(slide, "Text 42", "")


def drop_slide(prs, slide):
    """Remove a slide, located by part identity rather than by position: by
    the time the removals run, build_line_items_tables.build() has inserted
    the extra holdings pages and everything after slide 10 has moved."""
    line_items_mod.delete_slide(prs, [s.part for s in prs.slides].index(slide.part))


def fill_liquidity(prs, parsed):
    slide = get_slide(prs, 20)
    update_donut_by_name(slide, "Chart 5", parsed["liquidity_profile_pct"])
    profile = parsed["liquidity_profile_pct"]
    illiquid = profile.get("Illiquid (lock-up)", 0)
    daily = profile.get("Listed & daily-liquid", 0)
    set_text(slide, "TextBox 8",
             f"About {illiquid:.1f}% of the portfolio sits in private, hedge-fund and lock-up vehicles "
             f"with redemption gates and notice periods.\n\n"
             f"Liquidity events should be planned around these constraints; the {daily:.1f}% in listed "
             f"securities and daily-liquid funds covers near-term needs.")


def fill_concentration(prs, parsed):
    slide = get_slide(prs, 21)
    conc = parsed["concentration"]
    fill_concentration_table(slide, "Table 3", conc["top_holdings"])
    set_text(slide, "Rectangle 4", f"{conc['top_5_pct']:.1f}%\nTop 5")
    set_text(slide, "Rectangle 5", f"{conc['top_10_pct']:.1f}%\nTop 10")
    set_text(slide, "Rectangle 6", f"{conc['top_20_pct']:.1f}%\nTop 20")
    set_text(slide, "Rectangle 7", f"{conc['positions_above_5pct']}\nPositions > 5%")


def fill_income(prs, parsed):
    slide = get_slide(prs, 22)
    income = parsed["income"]
    ry = income["running_yield_pct"]
    set_text(slide, "Rectangle 5", f"{ry:.1f}%\nRunning yield" if ry is not None else "n/a\nRunning yield")
    fd = income["fi_duration_years"]
    set_text(slide, "Rectangle 6", f"{fd:.1f}y\nFI duration" if fd is not None else "n/a\nFI duration")
    ri = income["rate_impact_100bp_eur"]
    ccy = parsed["base_currency"]
    set_text(slide, "Rectangle 7", f"{ccy} {ri:,.0f}\nImpact +100bp" if ri is not None else "n/a\nImpact +100bp")
    fill_income_tables(slide, "Table 4", "Table 9", income)


# ------------------------------------------------------------- market slides
# Slides 5-6 (economic scenario / investment views) are permanently frozen
# per explicit instruction — they are never touched here, always shipping
# with the reference deck's own text. Only slides 3-4 pull from
# market_update.json.

def fill_market_slides(prs, market):
    if not market:
        return
    s3 = get_slide(prs, 3)
    h = market["headline"]
    # The template's own titles describe a weekly market recap ("How markets
    # moved last week"). A source that isn't one — a monthly house-view
    # summary, say — carries its own title in the JSON rather than shipping
    # under a heading its content doesn't match.
    if h.get("slide_title"):
        set_text(s3, "Title 1", h["slide_title"])
    set_subtitle(s3, h["intro_sentence"])

    rects = [sh for sh in s3.shapes if sh.name == "Rectangle"]
    stat_rects, obs_rects = rects[:4], rects[4:7]
    for sh, stat in zip(stat_rects, h["stats"]):
        set_shape_lines(sh, [stat["value"], stat["label"]])
    for sh, obs in zip(obs_rects, market["three_observations"]):
        set_shape_lines(sh, [obs["headline"], obs["body"]])

    textboxes = [sh for sh in s3.shapes if sh.name == "TextBox"]
    # [0]='Cross-asset scoreboard' label, [1]=Index col, [2]=Week col,
    # [3]=YTD col, [4]='Three observations' label
    if len(textboxes) >= 4:
        board = market.get("scoreboard") or {}
        rows = board.get("rows") or []
        # The block's label and its three column headers are data too, not
        # fixture: a source with no cross-asset performance table can fill the
        # same three columns with what it does have (e.g. a tactical-allocation
        # preference matrix: house view / stance / change). Both default to the
        # template's own weekly wording.
        headers = board.get("headers") or ["Index", "Week", "YTD"]
        if board.get("label"):
            set_shape_lines(textboxes[0], [board["label"]])
        if rows:
            fit_scoreboard_columns(
                textboxes[0], textboxes[1:4],
                [[headers[0]] + [r["index"] for r in rows],
                 [headers[1]] + [r["week"] for r in rows],
                 [headers[2]] + [r["ytd"] for r in rows]])
            set_shape_lines(textboxes[1], [headers[0]] + [r["index"] for r in rows])
            set_shape_lines(textboxes[2], [headers[1]] + [r["week"] for r in rows])
            set_shape_lines(textboxes[3], [headers[2]] + [r["ytd"] for r in rows])
            if board.get("colorize", True):
                color_scoreboard_column(textboxes[2], [r["week"] for r in rows])
                color_scoreboard_column(textboxes[3], [r["ytd"] for r in rows])
            else:
                # Sign-based green/orange is meaningless for non-numeric values,
                # and the template's own paragraphs carry the reference week's
                # colors — reset both columns to navy so a stance word doesn't
                # inherit "this number was negative".
                recolor_column(textboxes[2], NAVY)
                recolor_column(textboxes[3], NAVY)
        else:
            # No scoreboard in this source at all. "Leave it un-filled"
            # (SKILL.md step 3) has to mean *emptied* here: the template ships
            # with the reference week's real numbers in these boxes, so leaving
            # them alone would put another week's market data on the slide.
            for tb in textboxes[:4]:
                set_shape_lines(tb, [""])
            print("Note: market_update.json carries no scoreboard rows — slide 3's "
                  "scoreboard block was emptied rather than left with the "
                  "template's own reference-week numbers.")

    s4 = get_slide(prs, 4)
    if market["four_drivers"].get("slide_title"):
        set_text(s4, "Title 1", market["four_drivers"]["slide_title"])
    set_subtitle(s4, market["four_drivers"]["intro_sentence"])
    s4_rects = [sh for sh in s4.shapes if sh.name == "Rectangle"]
    for sh, drv in zip(s4_rects[:4], market["four_drivers"]["drivers"]):
        lines = [f"{drv['number']} - {drv['title']}", drv["body"]]
        if drv.get("stat"):
            lines.append(drv["stat"])
        # only the title (line 0) stays bold; body/stat lines forced normal
        set_shape_lines(sh, lines, bold_overrides={i: False for i in range(1, len(lines))})

    # The 5th 'Rectangle' is a standalone policy note (no title line), not a
    # numbered driver card — but it is still this week's market commentary.
    # Leaving it to the template ships the reference week's own policy
    # sentence to every client, so it is filled from the JSON, or emptied
    # when the source has no policy line to put there.
    if len(s4_rects) >= 5:
        note = market["four_drivers"].get("policy_note")
        set_shape_lines(s4_rects[4], [note or ""])
        if not note:
            print("Note: market_update.json has no four_drivers.policy_note — slide 4's "
                  "policy note was emptied rather than left with the template's own.")



# --------------------------------------------- private-equity monitoring ---

# The Alternatives sleeve slide (original template position 17) is the anchor:
# the fund-by-fund review goes immediately after it, inside the sleeve.
ALTERNATIVES_SLIDE = 17

# Column header, JSON key, and width in cm. Widths sum to the deck's own
# content width (26.11cm, the title placeholder's width).
PE_COLUMNS = [
    ("Fund",                   "name",                   4.00),
    ("Vintage",                "vintage",                1.50),
    ("Term",                   "term",                   1.30),
    ("Ccy",                    "currency",               1.00),
    ("As of",                  "as_of",                  1.80),
    ("TVPI / MOIC",            "multiple",               3.70),
    ("DPI",                    "dpi",                    1.30),
    ("Called",                 "capital_called",         1.50),
    ("Expected calls (4Q)",    "expected_calls",         4.00),
    ("Expected distrib. (4Q)", "expected_distributions", 6.01),
]
PE_TABLE_TOP = Emu(int(6.20 * 360000))
PE_ROW_HEIGHT = Emu(int(0.95 * 360000))
# The cards fill the band between the table and the page footnote (19.40cm),
# which leaves room for a title line plus six lines of 8pt commentary each.
PE_CARD_TOP = Emu(int(11.20 * 360000))
PE_CARD_WIDTH = Emu(int(12.85 * 360000))
PE_CARD_HEIGHT = Emu(int(3.85 * 360000))
PE_CARD_GAP_X = Emu(int(0.41 * 360000))
PE_CARD_GAP_Y = Emu(int(0.35 * 360000))
CONTENT_LEFT = Emu(665163)
CONTENT_WIDTH = Emu(9398001)


def _clone_decoration(shape, dest_slide):
    """Copy one decorative shape (the page footnote, the corner mark) onto a
    slide built from a bare layout.

    A picture also owns relationships, which do not travel with the XML: every
    r:embed in the copy still names an rId that only exists on the source
    slide, and PowerPoint reports the file as corrupt. Each one is re-related
    to the destination slide and rewritten. Note the corner mark is an SVG
    picture, so it carries *two* — `a:blip` for the raster fallback and
    `asvg:svgBlip` inside an extension for the vector original; walking every
    element with an r:embed catches both, where matching on `a:blip` alone
    silently leaves the second one dangling."""
    new_el = _copy.deepcopy(shape._element)
    dest_slide.shapes._spTree.append(new_el)
    embed_attr = qn("r:embed")
    remapped = {}
    for el in new_el.iter():
        old_rid = el.get(embed_attr)
        if not old_rid:
            continue
        if old_rid not in remapped:
            rel = shape.part.rels[old_rid]
            remapped[old_rid] = dest_slide.part.relate_to(rel.target_part, rel.reltype)
        el.set(embed_attr, remapped[old_rid])


def build_pe_monitoring_slide(prs, anchor_slide, pe):
    """Add the private-equity fund-by-fund review, immediately after the
    Alternatives sleeve slide it belongs to.

    Its content is a separate input (`--pe-monitoring`), not something the
    custodian Excel carries: manager-reported fund metrics — TVPI, DPI, called
    capital, expected calls and distributions — live in the managers' own
    quarterly reports, not in a custodian position file. Nothing here is
    derived from or reconciled against parsed.json.

    Like the slide removals in build(), this runs after the line-items
    rebuild, so the anchor is located by part identity rather than by a slide
    position the rebuild has already moved."""
    anchor_idx = [sl.part for sl in prs.slides].index(anchor_slide.part)
    slide = line_items_mod.duplicate_slide_after(prs, anchor_idx)

    # A slide built from the bare layout has the title and body placeholders
    # and nothing else — the footnote and corner mark live on each slide, so
    # they are copied from the anchor to keep the page furniture consistent.
    for name in ("TextBox 11", "Oval 12", "Graphic 20"):
        src = find_shape(anchor_slide, name)
        if src is not None:
            _clone_decoration(src, slide)
    # The title placeholder arrives empty (no runs), so its text goes in as a
    # new run and the font comes from the layout — the same inheritance every
    # other title on the deck uses. set_subtitle() then replaces the body
    # placeholder with the deck's standard fixed-position subtitle textbox.
    title = find_shape(slide, "Title 1")
    title.text_frame.paragraphs[0].add_run().text = pe.get(
        "slide_title", "Private equity: fund-by-fund review")
    set_subtitle(slide, pe["intro_sentence"])

    funds = pe["funds"]
    table_height = PE_ROW_HEIGHT * (len(funds) + 1)
    frame = slide.shapes.add_table(len(funds) + 1, len(PE_COLUMNS),
                                   CONTENT_LEFT, PE_TABLE_TOP, CONTENT_WIDTH, table_height)
    table = frame.table
    for c, (_, _, width_cm) in enumerate(PE_COLUMNS):
        table.columns[c].width = Emu(int(width_cm * 360000))
    for r in range(len(funds) + 1):
        table.rows[r].height = PE_ROW_HEIGHT
    for c, (header, _, _) in enumerate(PE_COLUMNS):
        _cell(table.cell(0, c), header, bold=True, fill=TABLE_HEADER_FILL,
              font_color=TABLE_HEADER_FONT, size=9)
    for r, fund in enumerate(funds, 1):
        for c, (_, key, _) in enumerate(PE_COLUMNS):
            _cell(table.cell(r, c), fund.get(key, "—"), size=9, bold=(c == 0))

    # One commentary card per fund, two by two under the table.
    for i, fund in enumerate(funds):
        left = CONTENT_LEFT + (PE_CARD_WIDTH + PE_CARD_GAP_X) * (i % 2)
        top = PE_CARD_TOP + (PE_CARD_HEIGHT + PE_CARD_GAP_Y) * (i // 2)
        box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), PE_CARD_WIDTH, PE_CARD_HEIGHT)
        box.name = f"PE Card {i + 1}"
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_FILL
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(144000)
        tf.margin_top = tf.margin_bottom = Emu(108000)
        for j, line in enumerate([fund["name"]] + list(fund["comments"])):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = line if j == 0 else f"- {line}"
            run.font.size = Pt(10 if j == 0 else 8)
            run.font.bold = (j == 0)
            run.font.color.rgb = NAVY
            para.space_after = Pt(3 if j == 0 else 2)
    return slide


# --------------------------------------------------------------------- main

def build(excel_path, profile, client_name, output_path, market_update_path=None,
          valuation_date=None, bucket_overrides=None, market_cap_overrides=None,
          sector_overrides=None, vehicle_overrides=None, keep_parsed_json=None,
          pe_monitoring_path=None):
    parse_cmd = [sys.executable, str(SCRIPT_DIR / "parse_portfolio.py"), excel_path,
                 "--profile", profile, "--client-name", client_name,
                 "-o", keep_parsed_json or str(SCRIPT_DIR.parent / "work" / "_parsed_tmp.json")]
    if valuation_date:
        parse_cmd += ["--valuation-date", valuation_date]
    if bucket_overrides:
        parse_cmd += ["--bucket-overrides", bucket_overrides]
    if market_cap_overrides:
        parse_cmd += ["--market-cap-overrides", market_cap_overrides]
    if sector_overrides:
        parse_cmd += ["--sector-overrides", sector_overrides]
    if vehicle_overrides:
        parse_cmd += ["--vehicle-overrides", vehicle_overrides]
    subprocess.run(parse_cmd, check=True)

    parsed_path = keep_parsed_json or str(SCRIPT_DIR.parent / "work" / "_parsed_tmp.json")
    parsed = json.loads(Path(parsed_path).read_text())

    market = None
    if market_update_path:
        market = json.loads(Path(market_update_path).read_text())

    pe_monitoring = None
    if pe_monitoring_path:
        pe_monitoring = json.loads(Path(pe_monitoring_path).read_text())

    prs = Presentation(str(TEMPLATE_PATH))

    # IMPORTANT: line_items_mod.build() inserts/deletes slides at position
    # 9-10, shifting every fixed slide number after it. Every other fill_*
    # function below addresses slides by their ORIGINAL template position
    # (get_slide(prs, N)), so it must run first, while slide 11 is still
    # slide 11 etc. line_items_mod.build() runs last, once every other
    # slide's content is already in place.
    fill_cover(prs, parsed)
    fill_market_slides(prs, market)
    fill_profile_dial(prs, parsed)
    fill_portfolio_overview(prs, parsed)
    fill_geographic_exposure(prs, parsed)
    fill_sleeve_slides(prs, parsed)
    fill_proposed_bond_selection(prs, parsed)
    fill_liquidity(prs, parsed)
    fill_concentration(prs, parsed)
    fill_income(prs, parsed)

    # Grab the slides addressed by their original template position now, while
    # those positions are still the template's; both operations below happen
    # after the line-items rebuild has moved everything after slide 10, and
    # find their slide by part identity instead.
    # The equity-breakdown slide always goes: its sector and market-cap donuts
    # are out of this deck by design, and its remaining panel (the equity
    # sleeve's regional split) has moved onto the geographic-exposure slide.
    doomed = [get_slide(prs, EQUITY_BREAKDOWN_SLIDE)]
    if not parsed.get("proposed_bond_selection"):
        doomed.append(get_slide(prs, PROPOSED_BOND_SLIDE))
        print("Note: the portfolio Excel has no 'Fixed Income' proposal tab — the "
              "'Fixed Income Breakdown' slide was removed rather than left showing "
              "the template's own example bond selection.")
    pe_anchor = get_slide(prs, ALTERNATIVES_SLIDE) if pe_monitoring else None
    line_items_mod.build(prs, parsed)  # slides 9-10(+): rebuilt as real tables, runs last
    for slide in doomed:
        drop_slide(prs, slide)
    if pe_anchor is not None:
        build_pe_monitoring_slide(prs, pe_anchor, pe_monitoring)
        print(f"Added the private-equity review slide ({len(pe_monitoring['funds'])} funds) "
              "after the Alternatives sleeve.")

    prs.save(output_path)
    print(f"Wrote {output_path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--excel", required=True)
    ap.add_argument("--profile", required=True,
                     choices=["Fixed Income", "Conservative", "Moderate", "Balanced", "Growth", "Equity"])
    ap.add_argument("--client-name", default="[Client Name]")
    ap.add_argument("--market-update", default=None, help="market_update.json produced from the market PDF")
    ap.add_argument("--pe-monitoring", default=None,
                     help="JSON of manager-reported private-market fund metrics; adds the "
                          "fund-by-fund review slide after the Alternatives sleeve (schema: "
                          "work/private_equity_monitoring_202608.json)")
    ap.add_argument("--valuation-date", default=None)
    ap.add_argument("--bucket-overrides", default=None)
    ap.add_argument("--market-cap-overrides", default=None)
    ap.add_argument("--sector-overrides", default=None)
    ap.add_argument("--vehicle-overrides", default=None,
                     help="CSV isin,vehicle — the desk's own Direct line / Fund / Structured "
                          "call, where the heuristic cannot tell them apart")
    ap.add_argument("--keep-parsed-json", default=None, help="also write parsed.json to this path")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()

    build(args.excel, args.profile, args.client_name, args.output,
          market_update_path=args.market_update, pe_monitoring_path=args.pe_monitoring,
          valuation_date=args.valuation_date,
          bucket_overrides=args.bucket_overrides, market_cap_overrides=args.market_cap_overrides,
          sector_overrides=args.sector_overrides, vehicle_overrides=args.vehicle_overrides,
          keep_parsed_json=args.keep_parsed_json)


if __name__ == "__main__":
    main()
