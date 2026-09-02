#!/usr/bin/env python3
"""
Helpers to push parsed.json data into the native PowerPoint charts already
sitting in a slide (donuts, clustered bars) — updates both the cached chart
values PowerPoint renders immediately and the embedded workbook it reopens
from, via python-pptx's chart.replace_data().

This module is imported by build_proposal.py; it is not meant to be run
standalone, but a CLI smoke-test is provided for quick checks:

    python update_chart.py deck.pptx --slide 11 --chart "Chart 10" \
        --data '{"Fixed Income": 55.4, "Equities": 26.3, "Cash": 5.4}'

Label conventions (assumptions.md §13):
- Native chart category labels (baked into the chart itself, as the
  reference deck does it) use a WHOLE-number percent: "Fixed Income  55%"
  (two spaces before the number).
- Companion textbox "legends" placed next to a donut (separate shapes, not
  part of the chart) use ONE-decimal percent: "United States   16.7%"
  (three spaces). Those are plain text boxes — see build_proposal.py's
  `set_text_by_name()`, not this module.
"""
from __future__ import annotations

import argparse
import json

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

# Brand palette (theme "Syz 2024") for donut slices the template does not
# color itself. Same values as build_proposal.py's constants, kept here as
# plain hex so this module stays importable on its own.
SLICE_PALETTE = ["3BAF90", "4B5F80", "AC5D85", "FFA400", "202945",
                 "FF6C0E", "FFC545", "79D6FF"]


def pct_label_whole(name: str, pct_value: float) -> str:
    # A slice big enough to be drawn but smaller than 0.5% would round to
    # "Cash  0%", which reads as an error rather than as "very small" — those
    # fall back to the one-decimal form (same two-space whole-label spacing).
    if round(pct_value) == 0:
        return f"{name}  " + ("<0.1%" if pct_value < 0.1 else f"{pct_value:.1f}%")
    return f"{name}  {round(pct_value):.0f}%"


def pct_label_one_decimal(name: str, pct_value: float) -> str:
    # a genuinely nonzero value that rounds to 0.0% reads as an error or
    # missing data next to a real category name, not as "very small"
    pct_str = "<0.1%" if 0 < pct_value < 0.1 else f"{pct_value:.1f}%"
    return f"{name}   {pct_str}"


def color_unstyled_slices(chart, n_slices: int) -> None:
    """Give an explicit brand color to every slice the template doesn't
    already color. Each donut in the deck carries a handful of hardcoded
    data-point colors (three on the liquidity donut, five on the asset-class
    one); slices past that are left to PowerPoint's own varied-color default,
    which lands outside the palette. Colors the template did specify are never
    touched — only the ones it left to chance get filled in."""
    series = chart.plots[0].series[0]
    styled = {int(dPt.find(qn("c:idx")).get("val")) for dPt in series._element.findall(qn("c:dPt"))}
    if styled >= set(range(n_slices)):
        return
    taken = {c.upper() for c in _explicit_colors(series)}
    spare = [c for c in SLICE_PALETTE if c not in taken] or SLICE_PALETTE
    for i in range(n_slices):
        if i in styled:
            continue
        fill = series.points[i].format.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(spare[i % len(spare)])


def _explicit_colors(series) -> list[str]:
    return [el.get("val") for el in series._element.iter(qn("a:srgbClr"))]


def find_chart_shape(slide, chart_name: str):
    for sh in slide.shapes:
        if sh.has_chart and sh.name == chart_name:
            return sh
    raise KeyError(f"No chart shape named {chart_name!r} on this slide "
                    f"(available: {[s.name for s in slide.shapes if s.has_chart]})")


def update_donut(chart_shape, data: dict[str, float], series_name: str = "", label_style: str = "whole"):
    """data: {category_name: pct_value_0_to_100}, already sorted the way you
    want slices to appear. Drops zero/near-zero slices (assumptions.md §5:
    a class with 0% is omitted, not shown as a sliver)."""
    label_fn = pct_label_whole if label_style == "whole" else pct_label_one_decimal
    items = [(k, v) for k, v in data.items() if v and v > 0.05]
    cd = CategoryChartData()
    cd.categories = [label_fn(k, v) for k, v in items]
    cd.add_series(series_name, [v / 100.0 for _, v in items])
    chart_shape.chart.replace_data(cd)
    color_unstyled_slices(chart_shape.chart, len(items))
    return chart_shape.chart


def update_bar(chart_shape, data: dict[str, float], series_name: str = ""):
    """Clustered bar charts: plain category names, no baked percentage
    (matches the reference deck's Geographic/Sector exposure bars)."""
    items = [(k, v) for k, v in data.items() if v and v > 0.0]
    cd = CategoryChartData()
    cd.categories = [k for k, _ in items]
    cd.add_series(series_name, [v / 100.0 for _, v in items])
    chart_shape.chart.replace_data(cd)
    return chart_shape.chart


def update_multi_series_bar(chart_shape, categories: list[str], series: dict[str, list[float]]):
    """For any bar/column chart with more than one series (none of the
    slides in slide_recipe.md currently need this, kept for extensibility)."""
    cd = CategoryChartData()
    cd.categories = categories
    for name, values in series.items():
        cd.add_series(name, values)
    chart_shape.chart.replace_data(cd)
    return chart_shape.chart


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx_path")
    ap.add_argument("--slide", type=int, required=True, help="1-indexed slide number")
    ap.add_argument("--chart", required=True, help="chart shape name, e.g. 'Chart 10'")
    ap.add_argument("--data", required=True, help='JSON object {"category": pct_value, ...}')
    ap.add_argument("--type", choices=["donut", "bar"], default="donut")
    ap.add_argument("-o", "--output", default=None)
    args = ap.parse_args()

    prs = Presentation(args.pptx_path)
    slide = list(prs.slides)[args.slide - 1]
    chart_shape = find_chart_shape(slide, args.chart)
    data = json.loads(args.data)
    if args.type == "donut":
        update_donut(chart_shape, data)
    else:
        update_bar(chart_shape, data)
    prs.save(args.output or args.pptx_path)
    print(f"Updated {args.chart!r} on slide {args.slide}")


if __name__ == "__main__":
    main()
